import openpyxl
from openpyxl.styles import Font
from docx import Document


def save_to_excel(lines, filename):
    """
    Save a list of instruction lines to an Excel (.xlsx) file.

    Args:
        lines (list of str): The lines to save.
        filename (str): The output Excel file path.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Instructions"

    bold_font = Font(bold=True)

    for idx, line in enumerate(lines, start=1):
        cell = ws.cell(row=idx, column=1, value=line)
        cell.font = bold_font if idx == 1 else None

    wb.save(filename)


def save_to_text(lines, filename):
    """
    Save a list of instruction lines to a plain text (.txt) file.

    Args:
        lines (list of str): The lines to save.
        filename (str): The output text file path.
    """
    with open(filename, "w", encoding="utf-8") as f:
        for line in lines:
            f.write(line + "\n")


def save_to_word(lines, filename):
    """
    Save a list of instruction lines to a Microsoft Word (.docx) file.

    Args:
        lines (list of str): The lines to save.
        filename (str): The output Word file path.
    """
    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    doc.save(filename)


def save_to_pdf(lines, filename):
    """
    Save a list of instruction lines to a PDF (.pdf) file.

    Args:
        lines (list of str): The lines to save.
        filename (str): The output PDF file path.
    """
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.units import inch
    
    # Create a PDF document with the letter page size
    doc = SimpleDocTemplate(filename, pagesize=letter)
    story = []
    styles = getSampleStyleSheet()
    
    # Add a title
    title_style = styles['Title']
    story.append(Paragraph("Installation Instructions", title_style))
    story.append(Spacer(1, 0.25 * inch))
    
    # Add each line as a paragraph with appropriate styling
    normal_style = styles['Normal']
    for line in lines:
        # Check if this is a section header (all caps)
        if line.isupper():
            # Use heading style for section headers
            heading_style = styles['Heading2']
            story.append(Spacer(1, 0.2 * inch))
            story.append(Paragraph(line, heading_style))
            story.append(Spacer(1, 0.1 * inch))
        else:
            # Use normal style for regular text
            story.append(Paragraph(line, normal_style))
            story.append(Spacer(1, 0.1 * inch))
    
    # Build the PDF document
    doc.build(story)
